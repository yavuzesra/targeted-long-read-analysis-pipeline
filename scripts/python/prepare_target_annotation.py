#!/usr/bin/env python3

"""
Create a standard gene-to-target-interval annotation table.

The input workbook contains gene coordinates and the corresponding custom
panel BED intervals. The output TSV uses BED 0-based, half-open coordinates.

One output row is produced for each gene-target interval relationship.
Consequently, one BED interval may occur more than once when it overlaps
more than one gene.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "Gene Symbol",
    "Custom panel BED chromosome(s)",
    "Custom panel BED start(s) 0-based",
    "Custom panel BED end(s) half-open",
}


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Prepare gene-to-target BED annotation."
    )
    parser.add_argument(
        "--workbook",
        required=True,
        type=Path,
        help="Annotation workbook path",
    )
    parser.add_argument(
        "--sheet",
        required=True,
        help="Worksheet name",
    )
    parser.add_argument(
        "--output",
        required=True,
        type=Path,
        help="Output TSV path",
    )
    return parser.parse_args()


def normalise_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def split_multivalue(value: Any) -> list[str]:
    text = normalise_text(value)

    if not text or text.upper() in {"NA", "N/A"}:
        return []

    return [
        item.strip()
        for item in re.split(r"\s*;\s*", text)
        if item.strip()
    ]


def parse_integer(value: str, label: str, gene: str) -> int:
    try:
        number = int(float(value))
    except ValueError as exc:
        raise ValueError(
            f"Gene '{gene}': invalid {label} value '{value}'."
        ) from exc

    return number


def read_annotation(
    workbook_path: Path,
    sheet_name: str,
) -> list[dict[str, str | int]]:
    if not workbook_path.is_file():
        raise ValueError(
            f"Annotation workbook not found: {workbook_path}"
        )

    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=True,
    )

    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet not found: {sheet_name}"
        )

    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError(
            f"Worksheet is empty: {sheet_name}"
        )

    headers = [
        normalise_text(value)
        for value in rows[0]
    ]

    missing = REQUIRED_COLUMNS - set(headers)

    if missing:
        raise ValueError(
            "Missing required annotation columns: "
            + ", ".join(sorted(missing))
        )

    column_index = {
        header: index
        for index, header in enumerate(headers)
        if header
    }

    optional_columns = {
        "match_status": "Custom panel match status",
        "coverage_percentage": "Custom panel gene coverage %",
        "missing_gene_bases": "Custom panel missing gene bases bp",
        "extra_gene_bases": "Custom panel extra bases outside gene bp",
    }

    records: list[dict[str, str | int]] = []

    for excel_row_number, row in enumerate(rows[1:], start=2):
        gene = normalise_text(
            row[column_index["Gene Symbol"]]
        )

        if not gene:
            continue

        chromosomes = split_multivalue(
            row[column_index["Custom panel BED chromosome(s)"]]
        )
        starts = split_multivalue(
            row[
                column_index[
                    "Custom panel BED start(s) 0-based"
                ]
            ]
        )
        ends = split_multivalue(
            row[
                column_index[
                    "Custom panel BED end(s) half-open"
                ]
            ]
        )

        if not chromosomes and not starts and not ends:
            continue

        lengths = {
            len(chromosomes),
            len(starts),
            len(ends),
        }

        if len(lengths) != 1:
            raise ValueError(
                f"Gene '{gene}' on Excel row {excel_row_number}: "
                "chromosome, start and end counts do not match."
            )

        for interval_number, (
            chromosome,
            start_text,
            end_text,
        ) in enumerate(
            zip(chromosomes, starts, ends),
            start=1,
        ):
            start = parse_integer(
                start_text,
                "BED start",
                gene,
            )
            end = parse_integer(
                end_text,
                "BED end",
                gene,
            )

            if start < 0 or end <= start:
                raise ValueError(
                    f"Gene '{gene}': invalid BED interval "
                    f"{chromosome}:{start}-{end}."
                )

            record: dict[str, str | int] = {
                "gene_symbol": gene,
                "gene_interval_number": interval_number,
                "chromosome": chromosome,
                "bed_start_0based": start,
                "bed_end_0based_exclusive": end,
                "genomic_start_1based": start + 1,
                "genomic_end_1based_inclusive": end,
                "target_length_bp": end - start,
                "target_key": (
                    f"{chromosome}:{start}-{end}"
                ),
                "source_excel_row": excel_row_number,
            }

            for output_name, workbook_column in (
                optional_columns.items()
            ):
                if workbook_column in column_index:
                    value = row[column_index[workbook_column]]
                    record[output_name] = normalise_text(value)
                else:
                    record[output_name] = ""

            records.append(record)

    if not records:
        raise ValueError(
            "No gene-to-target interval relationships were found."
        )

    return records


def validate_unique_gene_target_pairs(
    records: list[dict[str, str | int]],
) -> None:
    seen: set[tuple[str, str]] = set()

    for record in records:
        key = (
            str(record["gene_symbol"]),
            str(record["target_key"]),
        )

        if key in seen:
            raise ValueError(
                "Duplicate gene-target relationship detected: "
                f"{key[0]} / {key[1]}"
            )

        seen.add(key)


def write_annotation(
    records: list[dict[str, str | int]],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = [
        "gene_symbol",
        "gene_interval_number",
        "chromosome",
        "bed_start_0based",
        "bed_end_0based_exclusive",
        "genomic_start_1based",
        "genomic_end_1based_inclusive",
        "target_length_bp",
        "target_key",
        "match_status",
        "coverage_percentage",
        "missing_gene_bases",
        "extra_gene_bases",
        "source_excel_row",
    ]

    with output_path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
            delimiter="\t",
        )
        writer.writeheader()
        writer.writerows(records)


def main() -> int:
    args = parse_arguments()

    try:
        records = read_annotation(
            workbook_path=args.workbook,
            sheet_name=args.sheet,
        )

        validate_unique_gene_target_pairs(records)
        write_annotation(records, args.output)

    except ValueError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1

    genes = {
        str(record["gene_symbol"])
        for record in records
    }

    targets = {
        str(record["target_key"])
        for record in records
    }

    print(
        f"[INFO] Gene-target relationships: {len(records)}"
    )
    print(f"[INFO] Unique genes: {len(genes)}")
    print(f"[INFO] Unique target intervals: {len(targets)}")
    print(f"[INFO] Output: {args.output}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
