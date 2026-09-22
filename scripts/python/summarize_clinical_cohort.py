#!/usr/bin/env python3

"""
Create a combined small-variant summary for all enabled clinical samples.

Inputs are existing per-sample outputs:
- technical variant summary TSV
- complete EPI2ME ClinVar-annotated variant TSV
- review-candidate TSV

Outputs:
- one combined TSV
- one Excel workbook

This script reports EPI2ME and ClinVar annotations without performing
diagnostic or independent clinical classification.
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


MISSING_VALUES = {"", ".", "NA", "N/A", "na", "n/a"}

SUMMARY_COLUMNS = [
    "sample_id",
    "total_panel_variants",
    "pass_variants",
    "snvs",
    "insertions",
    "deletions",
    "complex_or_other",
    "heterozygous",
    "homozygous_alternate",
    "clinvar_annotated_variants",
    "review_candidates",
    "pathogenic_or_likely_pathogenic",
    "uncertain_significance",
    "conflicting_classifications",
    "benign_or_likely_benign",
    "other_clinvar_classifications",
]


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a combined clinical small-variant cohort summary."
    )
    parser.add_argument(
        "--samples",
        required=True,
        type=Path,
        help="Path to config/samples.tsv.",
    )
    parser.add_argument(
        "--outputs-root",
        required=True,
        type=Path,
        help="Clinical outputs directory, for example outputs/clinical.",
    )
    parser.add_argument(
        "--output-tsv",
        required=True,
        type=Path,
        help="Combined cohort summary TSV.",
    )
    parser.add_argument(
        "--output-xlsx",
        required=True,
        type=Path,
        help="Combined cohort summary Excel workbook.",
    )
    return parser.parse_args()


def is_missing(value: str | None) -> bool:
    return value is None or value.strip() in MISSING_VALUES


def load_enabled_clinical_samples(samples_path: Path) -> List[str]:
    if not samples_path.is_file():
        raise FileNotFoundError(
            f"Sample metadata file does not exist: {samples_path}"
        )

    sample_ids: List[str] = []

    with samples_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")

        required = {"sample_id", "analysis_role", "include"}

        if reader.fieldnames is None:
            raise ValueError("Sample metadata has no header.")

        missing_columns = required.difference(reader.fieldnames)

        if missing_columns:
            raise ValueError(
                "Sample metadata is missing required columns: "
                + ", ".join(sorted(missing_columns))
            )

        for row in reader:
            if (
                row["analysis_role"].strip()
                == "workflow_application"
                and row["include"].strip().lower() == "true"
            ):
                sample_ids.append(row["sample_id"].strip())

    if not sample_ids:
        raise ValueError(
            "No included workflow_application samples were found."
        )

    return sample_ids


def read_single_row_tsv(path: Path) -> Dict[str, str]:
    if not path.is_file():
        raise FileNotFoundError(f"Required summary file is missing: {path}")

    with path.open("r", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))

    if len(rows) != 1:
        raise ValueError(
            f"Expected exactly one data row in {path}, found {len(rows)}."
        )

    return rows[0]


def load_annotated_rows(path: Path) -> List[Dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(
            f"Required annotated variant table is missing: {path}"
        )

    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def count_data_rows(path: Path) -> int:
    if not path.is_file():
        raise FileNotFoundError(f"Required file is missing: {path}")

    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        next(reader, None)
        return sum(1 for row in reader if row)


def split_clinvar_terms(value: str) -> List[str]:
    """
    Split combined ClinVar labels while preserving their meaning.

    EPI2ME/ClinVar values may contain pipe-separated classifications.
    """
    if is_missing(value):
        return []

    return [
        term.strip().lower()
        for term in value.split("|")
        if term.strip()
    ]


def summarize_clinvar(
    annotated_rows: List[Dict[str, str]],
) -> Dict[str, int]:
    category_counts = Counter(
        {
            "pathogenic_or_likely_pathogenic": 0,
            "uncertain_significance": 0,
            "conflicting_classifications": 0,
            "benign_or_likely_benign": 0,
            "other_clinvar_classifications": 0,
        }
    )

    for row in annotated_rows:
        terms = split_clinvar_terms(
            row.get("clinvar_significance", "")
        )

        normalized = "|".join(terms)

        if "conflicting_classifications" in normalized:
            category_counts["conflicting_classifications"] += 1

        elif "uncertain_significance" in normalized:
            category_counts["uncertain_significance"] += 1

        elif (
            "pathogenic" in normalized
            or "likely_pathogenic" in normalized
        ):
            category_counts[
                "pathogenic_or_likely_pathogenic"
            ] += 1

        elif (
            "benign" in normalized
            or "likely_benign" in normalized
        ):
            category_counts["benign_or_likely_benign"] += 1

        else:
            category_counts[
                "other_clinvar_classifications"
            ] += 1

    return dict(category_counts)


def build_sample_summary(
    sample_id: str,
    outputs_root: Path,
) -> Dict[str, str | int]:
    sample_root = outputs_root / sample_id / "small_variants"

    technical_summary_path = (
        sample_root
        / "summary"
        / f"{sample_id}.variant_summary.tsv"
    )

    annotated_path = (
        sample_root
        / "variants"
        / f"{sample_id}.all_annotated_variants.tsv"
    )

    review_path = (
        sample_root
        / "variants"
        / f"{sample_id}.review_candidates.tsv"
    )

    technical = read_single_row_tsv(technical_summary_path)
    annotated_rows = load_annotated_rows(annotated_path)
    clinvar_counts = summarize_clinvar(annotated_rows)

    return {
        "sample_id": sample_id,
        "total_panel_variants": technical["total_variants"],
        "pass_variants": technical["pass_variants"],
        "snvs": technical["snvs"],
        "insertions": technical["insertions"],
        "deletions": technical["deletions"],
        "complex_or_other": technical["complex_or_other"],
        "heterozygous": technical["heterozygous"],
        "homozygous_alternate": technical[
            "homozygous_alternate"
        ],
        "clinvar_annotated_variants": len(annotated_rows),
        "review_candidates": count_data_rows(review_path),
        **clinvar_counts,
    }


def write_tsv(
    rows: List[Dict[str, str | int]],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=SUMMARY_COLUMNS,
            delimiter="\t",
        )
        writer.writeheader()
        writer.writerows(rows)


def write_excel(
    rows: List[Dict[str, str | int]],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Clinical cohort summary"

    worksheet.append(SUMMARY_COLUMNS)

    for row in rows:
        worksheet.append(
            [row[column] for column in SUMMARY_COLUMNS]
        )

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_index, column_name in enumerate(
        SUMMARY_COLUMNS,
        start=1,
    ):
        values = [
            str(worksheet.cell(row=row_index, column=column_index).value or "")
            for row_index in range(1, worksheet.max_row + 1)
        ]

        width = min(
            max(len(value) for value in values) + 2,
            35,
        )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width

    workbook.save(output_path)


def main() -> int:
    args = parse_arguments()

    samples_path = args.samples.expanduser().resolve()
    outputs_root = args.outputs_root.expanduser().resolve()
    output_tsv = args.output_tsv.expanduser().resolve()
    output_xlsx = args.output_xlsx.expanduser().resolve()

    try:
        sample_ids = load_enabled_clinical_samples(samples_path)

        rows = [
            build_sample_summary(sample_id, outputs_root)
            for sample_id in sample_ids
        ]

        write_tsv(rows, output_tsv)
        write_excel(rows, output_xlsx)

    except (
        FileNotFoundError,
        ValueError,
        KeyError,
    ) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print(f"Clinical samples summarized: {len(rows)}")
    print(f"TSV output:   {output_tsv}")
    print(f"Excel output: {output_xlsx}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
